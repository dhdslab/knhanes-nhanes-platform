# -*- coding: utf-8 -*-
"""Smoke tests for the evidence-RAG framework (no GPU / no real data)."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import evidence_rag as er  # noqa: E402
from evidence_rag import BM25, Chunk, EvidenceBase, EvidenceRAGAgent  # noqa: E402


def _chunks():
    return [
        Chunk("a", "In NHANES, Serum uric acid and Chronic kidney disease: odds ratio 1.43 (95% CI 1.35 to 1.52).",
              "manifest row 1", {"survey": "NHANES"}),
        Chunk("b", "In KNHANES, Alanine aminotransferase and metabolic syndrome: odds ratio 1.88 (95% CI 1.77 to 2.01).",
              "manifest row 2", {"survey": "KNHANES"}),
        Chunk("c", "In NHANES, Body mass index and metabolic syndrome: odds ratio 3.10 (95% CI 2.80 to 3.43).",
              "manifest row 3", {"survey": "NHANES"}),
    ]


def test_bm25_ranks_relevant_first():
    idx = BM25().fit(_chunks())
    hits = idx.search("uric acid chronic kidney disease", k=2)
    assert hits, "expected at least one hit"
    assert "uric acid" in hits[0][0].text.lower()


def test_bm25_empty_on_no_match():
    idx = BM25().fit(_chunks())
    assert idx.search("myocardial infarction troponin", k=3) == []


def test_parse_manifest_csv(tmp_path):
    p = tmp_path / "results.csv"
    p.write_text(
        "survey,exposure,outcome,measure,estimate,ci_low,ci_high,p_value,n_analytic,adjustment,fdr_q\n"
        "NHANES,Serum uric acid,Chronic kidney disease,OR,1.43,1.35,1.52,1e-22,26551,adjusted,4e-22\n",
        encoding="utf-8")
    chunks = er.parse_manifest_csv(str(p))
    assert len(chunks) == 1
    assert "odds ratio 1.43" in chunks[0].text
    assert chunks[0].meta["outcome"] == "Chronic kidney disease"


def test_parse_docx_roundtrip(tmp_path):
    from docx import Document
    d = Document()
    d.add_heading("Uric acid and CKD", level=1)
    d.add_paragraph("In NHANES the odds ratio was 1.43.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Estimate"
    t.rows[0].cells[1].text = "1.43"
    fp = tmp_path / "report.docx"
    d.save(str(fp))
    chunks = er.parse_docx(str(fp))
    joined = " ".join(c.text for c in chunks)
    assert "1.43" in joined
    assert "Uric acid and CKD" in joined


def test_evidence_agent_fallback_cites_sources():
    agent = EvidenceRAGAgent(enabled=False)
    idx = BM25().fit(_chunks())
    ans = agent.answer("uric acid and chronic kidney disease", idx, k=2)
    assert ans.source == "fallback"
    assert ans.citations
    assert "1.43" in ans.answer


def test_evidence_agent_reports_no_evidence():
    agent = EvidenceRAGAgent(enabled=False)
    idx = BM25().fit(_chunks())
    ans = agent.answer("relationship between troponin and stroke", idx, k=3)
    assert ans.citations == []
    assert "no direct evidence" in ans.answer.lower()


def test_evidence_base_ask_offline():
    base = EvidenceBase(_chunks(), use_llm=False)
    ans = base.ask("what predicts metabolic syndrome?", k=2)
    assert ans.citations
    assert any("metabolic syndrome" in c.snippet.lower() for c in ans.citations)
