# -*- coding: utf-8 -*-
"""Re-run a sample of the released corpus from the raw data and compare it with suppl/.

    python check_reproduction.py                 # all three checks
    python check_reproduction.py trend ml        # a subset

Needs the raw survey files under data/ and R. Output goes to a temporary SUPPL_DIR, so the
released corpus is never touched. Compared:

  trend   NHANES hypertension and KNHANES diabetes: every table cell of the trend report
  ml      NHANES high triglycerides: selected model, held-out AUROC and AUPRC, feature count,
          against _ml_regen_summary.csv
  assoc   NHANES hypertension: estimate, 95% CI, P, N and adjustment set of every exposure,
          against _manifest_association_NHANES.csv

Exits 0 when every compared value is identical.
"""
import os
import sys
import tempfile
import time

HERE = os.path.dirname(os.path.abspath(__file__))
SCR = tempfile.mkdtemp(prefix="repro_")
os.environ["SUPPL_DIR"] = os.path.join(SCR, "suppl")
sys.path.insert(0, HERE)
os.chdir(HERE)

import docx                                                          # noqa: E402
import numpy as np                                                   # noqa: E402
import pandas as pd                                                  # noqa: E402

import factory_core as fc                                            # noqa: E402
import ml_report                                                     # noqa: E402
import suppl_generator as sg                                         # noqa: E402
import trend_report                                                  # noqa: E402

REL = os.path.join(HERE, "suppl")


def tables(doc):
    return [[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]


def check_trend(ds, outcome):
    t0 = time.time()
    ddir, cyc = sg.DATA[ds]
    new = trend_report.build_trend_report(ds, ddir, list(cyc), outcome,
                                          {**sg.DEFS, "pop": {"age_min": max(sg.AGE, 19)}},
                                          "", "", False)
    p = os.path.join(SCR, f"trend_{ds}_{outcome}.docx")
    with open(p, "wb") as f:
        f.write(new)
    a = tables(docx.Document(p))
    b = tables(docx.Document(os.path.join(REL, ds, "trend", f"{sg.slug(outcome)}__trend_{ds}.docx")))
    cells = sum(len(r) for t in b for r in t)
    same = a == b
    print(f"trend  {ds:<7} {outcome:<7} {len(b)} tables, {cells} cells: "
          f"{'IDENTICAL' if same else 'DIFFERENT'}  ({time.time()-t0:.0f}s)", flush=True)
    return same


def check_ml(ds, outcome):
    t0 = time.time()
    ddir, cyc = sg.DATA[ds]
    s = {}
    ml_report.build_ml_report(ds, fc.load_raw(ds, ddir, cyc), outcome, fc.ml_features(ds, outcome),
                              sg.DEFS, "", "", False, models=sg.ML_MODELS,
                              sample_n=sg.ML_SAMPLE_N, summary=s)
    s = s["ml"]
    R = pd.read_csv(os.path.join(HERE, "_ml_regen_summary.csv"))
    r = R[(R.survey == ds) & (R.outcome == outcome)].iloc[0]
    ok = (s["best"] == r.best and round(s["auROC"], 4) == r.auROC
          and round(s["auPRC"], 4) == r.auPRC and s["n_feats"] == r.n_feats)
    print(f"ml     {ds:<7} {outcome:<7} {s['best']} AUROC {s['auROC']} AUPRC {s['auPRC']} "
          f"({s['n_feats']} features); released {r.best} {r.auROC} {r.auPRC} ({r.n_feats}): "
          f"{'IDENTICAL' if ok else 'DIFFERENT'}  ({time.time()-t0:.0f}s)", flush=True)
    return ok


def check_assoc(ds, outcome):
    t0 = time.time()
    ddir, cyc = sg.DATA[ds]
    d = fc.apply_definitions(fc.load_raw(ds, ddir, cyc), sg.DEFS)
    orig = fc.outcome_candidates
    fc.outcome_candidates = lambda _ds: [outcome]      # one outcome through the batch code path
    try:
        sg.gen_association(ds, d)
    finally:
        fc.outcome_candidates = orig
    new = pd.read_csv(os.path.join(sg.SUPPL, f"_manifest_association_{ds}.csv"))
    rel = pd.read_csv(os.path.join(REL, f"_manifest_association_{ds}.csv"))
    rel = rel[rel.out == outcome]
    m = rel.merge(new, on=["exp", "out"], suffixes=("_rel", "_new"), how="left")
    missing = int(m.est_new.isna().sum())
    diff = {c: float(np.nanmax(np.abs(m[c + "_rel"] - m[c + "_new"])))
            for c in ["est", "lo", "hi", "p", "n"]}
    adj = bool((m.adj_rel == m.adj_new).all())
    ok = missing == 0 and not any(diff.values()) and adj
    print(f"assoc  {ds:<7} {outcome:<7} {len(rel)} released rows, {missing} not reproduced, "
          f"max |difference| {diff}, adjustment sets equal: {adj}: "
          f"{'IDENTICAL' if ok else 'DIFFERENT'}  ({time.time()-t0:.0f}s)", flush=True)
    return ok


if __name__ == "__main__":
    what = sys.argv[1:] or ["trend", "ml", "assoc"]
    res = []
    if "trend" in what:
        res += [check_trend("NHANES", "htn"), check_trend("KNHANES", "dm")]
    if "ml" in what:
        res += [check_ml("NHANES", "high_tg")]
    if "assoc" in what:
        res += [check_assoc("NHANES", "htn")]
    print("ALL IDENTICAL" if all(res) else "SOME DIFFERENCES", flush=True)
    sys.exit(0 if all(res) else 1)
