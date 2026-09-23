# -*- coding: utf-8 -*-
"""Supplementary figures and tables describing the released corpus and its
relationship to the manuscript.

Everything here is computed from the two association manifests and the released
Word bundles, so each number can be traced to a file a reader has in hand.

Figures (written to Figures/), in the house style of figstyle.py
  SupplFigureS1  the whole evidence map as a volcano plot
  SupplFigureS2  per-outcome yield: exposures eligible and surviving FDR control
  SupplFigureS3  analytic sample size actually used, absolute and as a share
  SupplFigureS4  cross-survey concordance within each measurement family

Tables (written to _src/ as CSV, inserted into the Supplement by add_s11_s13.py)
  S11  contents and integrity of the released Word files
  S12  per-outcome corpus summary
"""
import os, sys, hashlib
import numpy as np, pandas as pd
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from scipy.stats import pearsonr

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                  # repository root: platform code and suppl/
BASE = os.path.join(HERE, "out")              # figures and merged bundles are written here
FIGS = os.path.join(BASE, "Figures")
os.makedirs(FIGS, exist_ok=True)
sys.path.insert(0, HERE)
sys.path.insert(0, ROOT)
import figstyle as fs
import factory_core as fc

K = pd.read_csv(os.path.join(ROOT, "suppl", "_manifest_association_KNHANES.csv"))
N = pd.read_csv(os.path.join(ROOT, "suppl", "_manifest_association_NHANES.csv"))
K["survey"], N["survey"] = "KNHANES", "NHANES"
A = pd.concat([K, N], ignore_index=True)
A["effect"] = np.where(A.measure == "OR", np.log(A.est.clip(lower=1e-12)), A.est)
A["sig"] = A.q < 0.05
A["nlq"] = -np.log10(A.q.clip(lower=1e-300))

# the eligible population and the participants with a usable survey design,
# from Reporting/participant_flow.csv (participant_flow.py)
ELIGIBLE = {"KNHANES": 110239, "NHANES": 47558}
DESIGN = {"KNHANES": 105756, "NHANES": 44249}


# ── S1: the evidence map ─────────────────────────────────────────────────────
def fig_s1():
    # A beta is in the units of its own outcome, so betas are not comparable across
    # outcomes and must not share an axis. For continuous outcomes the effect is
    # therefore shown as a partial correlation, r = t / sqrt(t^2 + n), derived from
    # the released estimate and confidence interval, which is unit-free.
    se = (A.hi - A.lo) / (2 * 1.96)
    t = A.est / se.replace(0, np.nan)
    A["partial_r"] = t / np.sqrt(t ** 2 + A.n)
    YCAP = 100                      # q values below 1e-100 are drawn at the cap

    fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.1), sharey=True)
    panels = [("OR", "effect", "log odds ratio per 1 SD", "Binary outcomes"),
              ("beta", "partial_r", "partial correlation", "Continuous outcomes")]
    for k, (ax, (meas, xcol, xlab, head)) in enumerate(zip(axes, panels)):
        d = A[(A.measure == meas) & A[xcol].notna()]
        counts = {}
        for sv in ("KNHANES", "NHANES"):
            s = d[d.survey == sv]
            counts[sv] = len(s)
            ax.scatter(s[xcol], s.nlq.clip(upper=YCAP), s=4.5, alpha=.38,
                       linewidths=0, color=fs.SURVEY_DK[sv], zorder=3)
        ax.axhline(-np.log10(0.05), ls=(0, (4, 2.5)), lw=0.9, color=fs.TEAL, zorder=4)
        ax.set_ylim(-4, YCAP + 8)
        ax.set_yticks([0, 25, 50, 75, 100])
        ax.set_xlabel(xlab)
        fs.panel(ax, grid="y")
        ax.annotate("FDR $q$ = 0.05", xy=(0.008, -np.log10(0.05)),
                    xycoords=("axes fraction", "data"), xytext=(0, 4),
                    textcoords="offset points", ha="left", va="bottom",
                    fontsize=7.2, color=fs.TEAL)
        ax.set_title(head, loc="left", fontsize=9, color=fs.INK, pad=16)
        ax.annotate("KNHANES $n$ = {:,}    NHANES $n$ = {:,}".format(
                        counts["KNHANES"], counts["NHANES"]),
                    xy=(0, 1.012), xycoords="axes fraction", ha="left", va="bottom",
                    fontsize=7.2, color=fs.INK2)
        fs.label(ax, "ab"[k], x=-0.175 if k == 0 else -0.055, y=1.135)
    axes[0].set_ylabel("$-$log$_{10}$ FDR $q$   (capped at 100)")
    fig.legend(handles=[Line2D([0], [0], marker="o", ls="none", markersize=4.6,
                               markerfacecolor=fs.SURVEY_DK[sv], markeredgecolor="none",
                               label=sv) for sv in ("KNHANES", "NHANES")],
               loc="lower right", bbox_to_anchor=(0.90, 0.955), ncol=2,
               handletextpad=0.35, columnspacing=1.6)
    fs.title(fig, "The released evidence map, all 7,478 association analyses")
    fig.subplots_adjust(wspace=0.12)
    return fs.save(fig, os.path.join(FIGS, "SupplFigureS1.png"))


# ── S2: per-outcome yield ────────────────────────────────────────────────────
def fig_s2():
    g = (A.groupby(["survey", "out"])
           .agg(analysed=("est", "size"), sig=("sig", "sum")).reset_index())
    order = g.groupby("out").analysed.sum().sort_values(ascending=False).index.tolist()[:34]
    labels = [fs.label_fit(fc.lab(o), 38) for o in order]

    fig, axes = plt.subplots(1, 2, figsize=(7.2, 6.4), sharey=True, sharex=True)
    y = np.arange(len(order))
    for k, (ax, sv) in enumerate(zip(axes, ("KNHANES", "NHANES"))):
        s = g[g.survey == sv].set_index("out").reindex(order)
        # eligible = an unfilled bar, survived = a filled one. A second mark rather
        # than a second hue, because the pale salmon and the neutral grey are not a
        # distinguishable pair (OKLab 8.8, deuteranopia 4.9).
        ax.barh(y, s.analysed.fillna(0), height=.68, facecolor="none",
                edgecolor=fs.GREY, lw=0.7, zorder=2)
        ax.barh(y, s.sig.fillna(0), height=.68, color=fs.SURVEY[sv],
                linewidth=0, zorder=3)
        ax.set_title(sv, loc="left", fontsize=9.5, color=fs.INK, fontweight="bold")
        fs.panel(ax, grid="x")
        ax.set_xlim(0, max(g.analysed) * 1.04)
        ax.set_ylim(len(order) - 0.45, -0.55)
    axes[0].set_yticks(y)
    axes[0].set_yticklabels(labels)
    axes[0].legend(handles=[Patch(facecolor="none", edgecolor=fs.GREY, lw=0.7,
                                  label="exposures eligible for testing"),
                            Patch(facecolor=fs.NAVY, label="survived FDR $q$ < 0.05")],
                   loc="lower left", bbox_to_anchor=(0, 1.045), ncol=2,
                   columnspacing=1.4)
    fs.title(fig, "Yield per outcome")
    fig.subplots_adjust(wspace=0.08)
    b0, b1 = axes[0].get_position(), axes[1].get_position()
    fig.text((b0.x0 + b1.x1) / 2, 0.055, "exposure–outcome pairs",
             ha="center", fontsize=9, color=fs.INK)
    for k, ax in enumerate(axes):
        bb = ax.get_position()
        fs.fig_label(fig, "ab"[k], bb.x0 - 0.028, bb.y1 + 0.012)
    return fs.save(fig, os.path.join(FIGS, "SupplFigureS2.png"))


# ── S3: the sample size actually used ────────────────────────────────────────
def fig_s3():
    """Absolute contributing n, and the same thing as a share of the survey's
    usable-design sample, which is what makes the shortfall legible: the eligible
    population is three times the largest analysis, so a reference line at it
    would leave two thirds of the panel empty."""
    fig, axes = plt.subplots(1, 2, figsize=(7.2, 2.9))

    ax = axes[0]
    bins = np.linspace(0, A.n.max() * 1.02, 46)
    for sv in ("KNHANES", "NHANES"):
        s = A[A.survey == sv].n
        ax.hist(s, bins=bins, color=fs.SURVEY[sv], alpha=.55, linewidth=0, zorder=3)
        ax.hist(s, bins=bins, histtype="step", color=fs.SURVEY_DK[sv], lw=0.9, zorder=4)
    ax.set_ylim(top=ax.get_ylim()[1] * 1.20)
    for sv, va, ha, dx in (("KNHANES", 0.99, "right", -4), ("NHANES", 0.74, "right", -4)):
        m = A[A.survey == sv].n.median()
        ax.axvline(m, ls=(0, (3, 2)), lw=0.9, color=fs.SURVEY_DK[sv], zorder=5)
        ax.annotate(sv + "\n" + f"median {m:,.0f}",
                    xy=(m, va), xycoords=("data", "axes fraction"),
                    xytext=(dx, 0), textcoords="offset points", ha=ha, va="top",
                    fontsize=7.2, color=fs.SURVEY_DK[sv], linespacing=1.35)
    ax.set_xlabel("participants contributing to the fitted model")
    ax.set_ylabel("analyses")
    ax.set_xlim(0, A.n.max() * 1.02)
    fs.panel(ax, grid="y"); fs.thousands(ax)
    ax.set_title("Absolute", loc="left", fontsize=9, color=fs.INK, pad=16)
    fs.label(ax, "a", x=-0.165, y=1.105)

    ax = axes[1]
    bins = np.linspace(0, 100, 46)
    for sv in ("KNHANES", "NHANES"):
        s = 100 * A[A.survey == sv].n / DESIGN[sv]
        ax.hist(s, bins=bins, color=fs.SURVEY[sv], alpha=.55, linewidth=0, zorder=3)
        ax.hist(s, bins=bins, histtype="step", color=fs.SURVEY_DK[sv], lw=0.9, zorder=4)
    ax.set_xlim(0, 100)
    ax.set_xlabel("share of the survey's usable-design sample (%)")
    fs.panel(ax, grid="y")
    ax.set_title("Relative", loc="left", fontsize=9, color=fs.INK, pad=16)
    ax.annotate("usable-design sample {:,} and {:,}".format(
                    DESIGN["KNHANES"], DESIGN["NHANES"]),
                xy=(0, 1.012), xycoords="axes fraction", ha="left", va="bottom",
                fontsize=7.2, color=fs.INK2)
    ax.legend(handles=[Line2D([0], [0], color=fs.SURVEY_DK[sv], lw=1.6, label=sv)
                       for sv in ("KNHANES", "NHANES")],
              loc="upper right", bbox_to_anchor=(1.0, 1.0), handlelength=1.3)
    fs.label(ax, "b", x=-0.085, y=1.105)

    fs.title(fig, "Sample size contributing to each fitted analysis")
    fig.subplots_adjust(wspace=0.26)
    return fs.save(fig, os.path.join(FIGS, "SupplFigureS3.png"))


# ── S4: concordance within measurement families ──────────────────────────────
def fig_s4():
    ko = K[(K.measure == "OR") & (K.est > 0)][["exp", "out", "est"]].rename(columns={"est": "eK"})
    no = N[(N.measure == "OR") & (N.est > 0)][["exp", "out", "est"]].rename(columns={"est": "eN"})
    M = ko.merge(no, on=["exp", "out"])
    M["lK"], M["lN"] = np.log(M.eK), np.log(M.eN)
    fam = {}
    for i, F in enumerate(fc._FAMILIES):
        for v in F:
            fam[v] = i
    names = {0: "smoking", 1: "body composition", 2: "lipids", 3: "liver",
             4: "glycaemia", 5: "kidney"}
    M["fam"] = [fam.get(e, -1) for e in M.exp]
    groups = [(-1, "other / unassigned")] + [(i, names.get(i, str(i))) for i in sorted(names)]
    groups = [(i, nm) for i, nm in groups if (M.fam == i).sum() >= 20]
    groups.sort(key=lambda g: -(M.fam == g[0]).sum())

    ncol = 4
    nrow = int(np.ceil(len(groups) / ncol))
    fig, axes = plt.subplots(nrow, ncol, figsize=(7.2, 1.92 * nrow))
    axes = np.atleast_1d(axes).ravel()
    for k, (ax, (i, nm)) in enumerate(zip(axes, groups)):
        s = M[M.fam == i]
        r = pearsonr(s.lK, s.lN)[0]
        lo = min(s.lK.min(), s.lN.min()); hi = max(s.lK.max(), s.lN.max())
        pad = (hi - lo) * 0.06
        ax.plot([lo - pad, hi + pad], [lo - pad, hi + pad], ls=(0, (4, 2.5)),
                lw=0.9, color=fs.TEAL, zorder=2)
        ax.scatter(s.lK, s.lN, s=5.5, alpha=.45, color=fs.NAVY, linewidths=0, zorder=3)
        ax.set_xlim(lo - pad, hi + pad); ax.set_ylim(lo - pad, hi + pad)
        ax.set_aspect("equal", adjustable="box")
        ax.set_title(nm, loc="left", fontsize=9, color=fs.INK)
        ax.annotate(f"$n$ = {len(s):,}\n$r$ = {r:.2f}", xy=(0.035, 0.965),
                    xycoords="axes fraction", va="top", ha="left",
                    fontsize=7.2, color=fs.INK2, linespacing=1.35)
        ax.locator_params = None
        ax.xaxis.set_major_locator(plt.MaxNLocator(4))
        ax.yaxis.set_major_locator(plt.MaxNLocator(4))
        fs.panel(ax, grid="both")
        fs.label(ax, "abcdefg"[k], x=-0.24, y=1.02, size=10)
        if k % ncol == 0:
            ax.set_ylabel("NHANES  log OR")
        if k >= len(groups) - ncol:
            ax.set_xlabel("KNHANES  log OR")
    for ax in axes[len(groups):]:
        ax.axis("off")
    # the last row is short, so its panels need their own x label
    for k in range(len(groups)):
        if k + ncol >= len(groups):
            axes[k].set_xlabel("KNHANES  log OR")
    fs.title(fig, "Cross-survey concordance within each exposure measurement family")
    fig.subplots_adjust(wspace=0.42, hspace=0.42)
    return fs.save(fig, os.path.join(FIGS, "SupplFigureS4.png"))


# ── S11: released file inventory ─────────────────────────────────────────────
def md5(p):
    h = hashlib.md5()
    with open(p, "rb") as f:
        for c in iter(lambda: f.read(1 << 20), b""):
            h.update(c)
    return h.hexdigest()


def table_s11():
    import docx
    spec = [("Supplementary_Appendix_fixed.docx", "both", "Methods appendix S1-S15", None),
            ("Supplementary_Data_S1.docx", "KNHANES", "association reports", 3522),
            ("Supplementary_Data_S2.docx", "NHANES", "association reports", 3956),
            ("Supplementary_Data_S3.docx", "KNHANES", "temporal-trend reports", 27),
            ("Supplementary_Data_S4.docx", "NHANES", "temporal-trend reports", 25),
            ("Supplementary_Data_S5.docx", "KNHANES", "machine-learning reports", 27),
            ("Supplementary_Data_S6.docx", "NHANES", "machine-learning reports", 27)]
    rows = []
    for fn, sv, kind, n in spec:
        p = os.path.join(BASE, "Suppl", fn)
        if not os.path.exists(p):
            continue
        d = docx.Document(p)
        rows.append(dict(file=fn.replace("_fixed", ""), survey=sv, contents=kind,
                         reports=("-" if n is None else f"{n:,}"),
                         tables=f"{len(d.tables):,}", size_mb=f"{os.path.getsize(p)/1e6:.1f}",
                         md5=md5(p)[:12]))
        print(f"  {fn}: {len(d.tables)} tables")
    if not rows:
        # the merged bundles are journal files, not part of the repository: keep the committed table
        print(f"  no bundles in {os.path.join(BASE, 'Suppl')}; S11_file_inventory.csv left as is")
        return pd.read_csv(os.path.join(HERE, "S11_file_inventory.csv"))
    T = pd.DataFrame(rows)
    T.to_csv(os.path.join(HERE, "S11_file_inventory.csv"), index=False, encoding="utf-8-sig")
    return T


# ── S12: per-outcome corpus summary ──────────────────────────────────────────
def table_s12():
    g = (A.groupby(["survey", "out", "measure"])
           .agg(pairs=("est", "size"), sig=("sig", "sum"),
                med_eff=("effect", lambda s: np.median(np.abs(s))),
                med_n=("n", "median"), min_n=("n", "min"), max_n=("n", "max"))
           .reset_index())
    g["outcome"] = [fc.lab(o) for o in g.out]
    g["pct_sig"] = 100 * g.sig / g.pairs
    g = g.sort_values(["survey", "pairs"], ascending=[True, False])
    out = g[["survey", "outcome", "measure", "pairs", "sig", "pct_sig",
             "med_eff", "med_n", "min_n", "max_n"]]
    out.to_csv(os.path.join(HERE, "S12_outcome_summary.csv"), index=False, encoding="utf-8-sig")
    return out


if __name__ == "__main__":
    print("figures:")
    for f in (fig_s1, fig_s2, fig_s3, fig_s4):
        print("  wrote", os.path.basename(f()))
    print("tables:")
    t11 = table_s11(); print(f"  S11: {len(t11)} files")
    t12 = table_s12(); print(f"  S12: {len(t12)} outcome-by-survey rows")
    print(f"\ncorpus: {len(A):,} association analyses, {int(A.sig.sum()):,} FDR-significant "
          f"({100*A.sig.mean():.1f}%)")
